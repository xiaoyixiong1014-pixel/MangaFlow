"""导出能力接口测试模块（文件级说明，Task 7）。

覆盖链路（数据以 ORM 直插为主，绕过 LLM 链路）：
- 项目 ZIP 导出：manifest 结构（version/project/episodes/scripts/characters/
  storyboards/storyboard_characters/missing_media）与 media/ 文件打包
- 项目 ZIP 导入：manifest 校验（缺失/版本不符 400）、全链重建（新 ID +
  外键重映射）、媒体文件落回（同名冲突加时间戳）、URL 修正
- 分镜表导出：xlsx（openpyxl 可解析、表头行）与 csv（utf-8-sig、含表头）
- 成片 MP4 下载：原片直下（字节一致）、resolution 档位请求（转码失败降级
  原片）、非本人成片 404
"""

import io
import json
import uuid
import zipfile
from pathlib import Path
from typing import Any, Dict

import httpx
import openpyxl
import pytest
from sqlalchemy import select

from app.core.config import settings
from app.models import (
    Character,
    Episode,
    FilmOutput,
    Script,
    Storyboard,
    StoryboardCharacter,
)
from tests.conftest import create_project, register_user, unique_suffix


def _storage_root() -> Path:
    """获取本地存储根目录（测试造媒体文件用，内部函数）。

    Returns:
        Path: settings.storage_path() 对应的绝对路径。

    Raises:
        无：纯配置读取。
    """
    return settings.storage_path()


async def _seed_full_project(client: httpx.AsyncClient, auth: Dict[str, Any],
                             db, tag: str) -> Dict[str, Any]:
    """辅助：造一个"全集"项目（剧集/剧本/角色/分镜/角色强绑/媒体文件）。

    媒体文件真实写入 storage/generated/（导出 ZIP 时会被打包进 media/），
    测试结束统一清理。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。
        db: 测试库会话。
        tag: 防冲突随机后缀。

    Returns:
        Dict[str, Any]: {"project": 项目 JSON, "episode": 剧集 ORM,
        "character": 角色 ORM, "storyboard": 分镜 ORM, "media_files": 待清理路径列表,
        "media_url": 媒体 /storage URL}。

    Raises:
        AssertionError: 建项目接口返回非 200 时中断。
    """
    project = await create_project(
        client, auth["headers"], title=f"导出项目_{tag}", style="日漫", target_duration=60
    )
    pid = project["id"]

    # 真实媒体文件（角色定妆照）
    media_name = f"export_test_{tag}.png"
    media_path = _storage_root() / "generated" / media_name
    media_path.parent.mkdir(parents=True, exist_ok=True)
    media_path.write_bytes(b"\x89PNG\r\n\x1a\nfake-image-bytes-" + tag.encode())

    episode = Episode(project_id=pid, title=f"第一集_{tag}", sequence=1, status="draft")
    db.add(episode)
    await db.flush()

    script = Script(
        project_id=pid,
        episode_id=episode.id,
        content={"title": {"value": "测试剧本", "source": "ai"}},
        raw_text="林小满：你好。",
        version=1,
    )
    character = Character(
        project_id=pid,
        name="林小满",
        visual_prompt="少女，短发，校服",
        asset_url=f"/storage/generated/{media_name}",
        face_locked=False,
        generation_status="done",
    )
    db.add_all([script, character])
    await db.flush()

    storyboard = Storyboard(
        project_id=pid,
        episode_id=episode.id,
        script_id=script.id,
        sequence=1,
        shot_type="中景",
        camera_movement="固定",
        duration=6,
        visual_description="天台上两人对话",
        dialogue="林小满：你好\n旁白：风起了",
        scene_id=None,
        prompt_locked="锁定提示词：统一风格",
        prompt_editable="",  # 留空 → 导出应回退到 prompt_locked
        keyframe_url=f"/storage/generated/{media_name}",
        status="done",
    )
    db.add(storyboard)
    await db.flush()
    db.add(StoryboardCharacter(
        storyboard_id=storyboard.id, character_id=character.id,
        variant_id=None, link_confidence=0.8,
    ))
    await db.commit()

    return {
        "project": project,
        "episode": episode,
        "character": character,
        "storyboard": storyboard,
        "media_files": [media_path],
        "media_url": f"/storage/generated/{media_name}",
    }


async def test_zip_export_manifest_and_media(client, auth, db):
    """ZIP 导出：manifest 各域数据齐全、media/ 打包本地文件、project 去 owner。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。
        db: 测试库会话。

    Returns:
        None。

    Raises:
        AssertionError: 任一断言失败时中断。
    """
    tag = unique_suffix()
    seeded = await _seed_full_project(client, auth, db, tag)
    pid = seeded["project"]["id"]

    resp = await client.get(f"/api/projects/{pid}/export", headers=auth["headers"])
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/zip")
    assert "attachment" in resp.headers["content-disposition"]

    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        names = zf.namelist()
        assert "manifest.json" in names
        manifest = json.loads(zf.read("manifest.json").decode("utf-8"))
        assert manifest["version"] == 1
        # project 全字段但去掉 owner（user_id）
        assert "user_id" not in manifest["project"]
        assert manifest["project"]["id"] == pid
        assert manifest["project"]["title"].startswith("导出项目_")
        # 各域数据量一致
        assert len(manifest["episodes"]) == 1
        assert len(manifest["scripts"]) == 1
        assert len(manifest["characters"]) == 1
        assert len(manifest["storyboards"]) == 1
        assert len(manifest["storyboard_characters"]) == 1
        # 分镜字段完整（对白/提示词锁定层随行导出）
        assert manifest["storyboards"][0]["dialogue"] == "林小满：你好\n旁白：风起了"
        assert manifest["storyboards"][0]["prompt_locked"] == "锁定提示词：统一风格"
        # 媒体文件按 URL 相对路径打包进 media/
        rel = seeded["media_url"].removeprefix("/storage/")
        assert f"media/{rel}" in names
        assert zf.read(f"media/{rel}") == b"\x89PNG\r\n\x1a\nfake-image-bytes-" + tag.encode()

    # 清理造的媒体文件
    for p in seeded["media_files"]:
        p.unlink(missing_ok=True)


async def test_zip_import_rebuild_and_media_remap(client, auth, db):
    """ZIP 导入：全链重建（新项目/新外键）、媒体落回并修正 URL、规模断言。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。
        db: 测试库会话。

    Returns:
        None。

    Raises:
        AssertionError: 任一断言失败时中断。
    """
    tag = unique_suffix()
    seeded = await _seed_full_project(client, auth, db, tag)
    old_pid = seeded["project"]["id"]
    old_media_rel = seeded["media_url"].removeprefix("/storage/")

    export = await client.get(f"/api/projects/{old_pid}/export", headers=auth["headers"])
    assert export.status_code == 200

    # 导入（multipart 上传导出的 ZIP）
    imported = await client.post(
        "/api/projects/import",
        headers=auth["headers"],
        files={"file": ("project.zip", export.content, "application/zip")},
    )
    assert imported.status_code == 200, imported.text
    body = imported.json()
    new_pid = body["project_id"]
    assert new_pid != old_pid
    assert body["episodes"] == 1
    assert body["storyboards"] == 1
    assert body["storyboard_characters"] == 1

    try:
        # 新项目详情：剧集/剧本计数一致
        detail = await client.get(f"/api/projects/{new_pid}", headers=auth["headers"])
        assert detail.status_code == 200, detail.text
        assert detail.json()["episode_count"] == 1
        assert detail.json()["script_count"] == 1
        assert detail.json()["title"].endswith("（导入）")

        # 分镜全链重建：数量一致且角色强绑（character_links）重映射成功
        sbs = await client.get(f"/api/projects/{new_pid}/storyboards", headers=auth["headers"])
        assert sbs.status_code == 200, sbs.text
        assert len(sbs.json()) == 1
        assert sbs.json()[0]["characters"][0]["name"] == "林小满"

        # 分镜归属新剧集（外键重映射，从库内断言）
        new_ep = (await db.execute(
            select(Episode).where(Episode.project_id == new_pid)
        )).scalar_one()
        new_sb = (await db.execute(
            select(Storyboard).where(Storyboard.project_id == new_pid)
        )).scalar_one()
        assert new_sb.episode_id == new_ep.id

        # 媒体落回：同名冲突 → 文件名加时间戳，URL 已修正且物理文件存在
        new_char = (await db.execute(
            select(Character).where(Character.project_id == new_pid)
        )).scalar_one()
        assert new_char.asset_url != seeded["media_url"]  # 冲突改名
        assert new_char.asset_url.startswith("/storage/generated/")
        new_rel = new_char.asset_url.removeprefix("/storage/")
        assert new_rel != old_media_rel
        assert (_storage_root() / new_rel).is_file()
        # 落回文件内容与导出的原文件一致
        assert (_storage_root() / new_rel).read_bytes().endswith(tag.encode())

        # 新分镜的关键帧 URL 同步修正到新媒体文件
        assert new_sb.keyframe_url == new_char.asset_url
    finally:
        # 清理导入落盘的新媒体文件与造的原文件
        for rel in {old_media_rel, new_rel or ""}:
            if rel:
                (_storage_root() / rel).unlink(missing_ok=True)


async def test_zip_import_invalid_manifest_400(client, auth):
    """ZIP 导入校验：缺 manifest / 版本不符 / 非 ZIP 均返回 400。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。

    Returns:
        None。

    Raises:
        AssertionError: 任一断言失败时中断。
    """
    # 非 ZIP 内容
    resp = await client.post(
        "/api/projects/import",
        headers=auth["headers"],
        files={"file": ("bad.zip", b"not-a-zip", "application/zip")},
    )
    assert resp.status_code == 400

    # 缺 manifest.json 的合法 ZIP
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("other.txt", "hello")
    resp = await client.post(
        "/api/projects/import",
        headers=auth["headers"],
        files={"file": ("nomanifest.zip", buf.getvalue(), "application/zip")},
    )
    assert resp.status_code == 400

    # 版本不符
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("manifest.json", json.dumps({"version": 999, "project": {}}))
    resp = await client.post(
        "/api/projects/import",
        headers=auth["headers"],
        files={"file": ("wrongver.zip", buf.getvalue(), "application/zip")},
    )
    assert resp.status_code == 400


async def test_storyboard_export_csv_and_xlsx(client, auth, db):
    """分镜表导出：csv 含表头（utf-8-sig）、xlsx 可被 openpyxl 解析且表头行正确。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。
        db: 测试库会话。

    Returns:
        None。

    Raises:
        AssertionError: 任一断言失败时中断。
    """
    tag = unique_suffix()
    project = await create_project(client, auth["headers"], title=f"分镜导出_{tag}")
    pid = project["id"]
    episode = Episode(project_id=pid, title=f"EP01_{tag}", sequence=1, status="draft")
    db.add(episode)
    await db.flush()
    for seq, (shot, desc, dlg) in enumerate(
        [("中景", "天台对话", "林小满：你好\n旁白：风起了"),
         ("远景", "城市夜景", "旁白：夜色渐深")], start=1,
    ):
        db.add(Storyboard(
            project_id=pid, episode_id=episode.id, sequence=seq,
            shot_type=shot, duration=5, visual_description=desc, dialogue=dlg,
            prompt_locked=f"锁定提示词_{seq}", prompt_editable="",
            status="pending",
        ))
    await db.commit()

    base = f"/api/projects/{pid}/storyboards/episodes/{episode.id}/export"

    # ---- CSV：utf-8-sig、含表头与数据 ----
    csv_resp = await client.get(base, params={"format": "csv"}, headers=auth["headers"])
    assert csv_resp.status_code == 200, csv_resp.text
    assert csv_resp.headers["content-type"].startswith("text/csv")
    text = csv_resp.content.decode("utf-8-sig")
    for header in ("编号", "场景", "景别", "画面描述", "台词", "旁白/口播",
                   "预估时长", "提示词", "状态", "关键帧 URL"):
        assert header in text
    assert "天台对话" in text
    assert "锁定提示词_1" in text  # prompt_editable 为空 → 回退 prompt_locked
    assert "夜色渐深" in text

    # ---- XLSX：MIME 正确、openpyxl 可解析、表头行 + 2 行数据 ----
    xlsx_resp = await client.get(base, params={"format": "xlsx"}, headers=auth["headers"])
    assert xlsx_resp.status_code == 200, xlsx_resp.text
    assert xlsx_resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    ws = wb.active
    assert ws.title == "分镜表"
    assert ws.max_row == 3  # 表头 + 2 镜
    assert [c.value for c in ws[1]][:3] == ["编号", "场景", "景别"]
    assert ws.cell(row=2, column=4).value == "天台对话"
    assert ws.cell(row=2, column=7).value == 5

    # ---- 非法 format → 400（SB-08 起支持 pdf；非法值由实现手动校验返回 400） ----
    bad = await client.get(base, params={"format": "docx"}, headers=auth["headers"])
    assert bad.status_code == 400


async def test_storyboard_export_all_with_episode_zero(client, auth, db):
    """分镜表导出全量模式：episode_id=0 导出项目全部分镜（含无剧集镜头，Task 10.4）。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。
        db: 测试库会话。

    Returns:
        None。

    Raises:
        AssertionError: 任一断言失败时中断。
    """
    tag = unique_suffix()
    project = await create_project(client, auth["headers"], title=f"全量导出_{tag}")
    pid = project["id"]
    # 两条无剧集分镜（普通创作流拆解产物，episode_id 为空）
    for seq in (1, 2):
        db.add(Storyboard(
            project_id=pid, sequence=seq, shot_type="中景", duration=5,
            visual_description=f"全量导出画面_{seq}", dialogue="",
            prompt_locked=f"锁定_{seq}", status="pending",
        ))
    await db.commit()

    resp = await client.get(
        f"/api/projects/{pid}/storyboards/episodes/0/export",
        params={"format": "csv"}, headers=auth["headers"],
    )
    assert resp.status_code == 200, resp.text
    text = resp.content.decode("utf-8-sig")
    assert "全量导出画面_1" in text
    assert "全量导出画面_2" in text
    # 文件名为全量语义
    assert "storyboards_all" in resp.headers.get("content-disposition", "")


async def test_storyboard_export_denied_for_other_user(client, auth, db):
    """分镜表导出归属校验：他人项目的剧集导出 404。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。
        db: 测试库会话。

    Returns:
        None。

    Raises:
        AssertionError: 任一断言失败时中断。
    """
    tag = unique_suffix()
    project = await create_project(client, auth["headers"], title=f"他人项目_{tag}")
    episode = Episode(project_id=project["id"], title=f"EP_{tag}", sequence=1, status="draft")
    db.add(episode)
    await db.commit()

    other = await register_user(client, prefix="other")
    other_headers = {"Authorization": f"Bearer {other['access_token']}"}
    resp = await client.get(
        f"/api/projects/{project['id']}/storyboards/episodes/{episode.id}/export",
        params={"format": "csv"}, headers=other_headers,
    )
    assert resp.status_code == 404


async def test_film_output_download_original_and_resolution(client, auth, db):
    """成片下载：原片直下字节一致；带 resolution 档位请求仍 200（转码失败
    自动降级原片）；非本人成片 404；非法 resolution 400。

    Args:
        client: 内存 HTTP 客户端。
        auth: 随机注册用户上下文。
        db: 测试库会话。

    Returns:
        None。

    Raises:
        AssertionError: 任一断言失败时中断。
    """
    tag = unique_suffix()
    project = await create_project(
        client, auth["headers"], title=f"下载项目_{tag}", aspect_ratio="9:16"
    )
    pid = project["id"]

    # 造一个本地"成片"文件（内容非真实 MP4：ffmpeg 转码必失败 → 验证降级路径）
    film_name = f"download_test_{tag}.mp4"
    film_path = _storage_root() / "outputs" / film_name
    film_path.parent.mkdir(parents=True, exist_ok=True)
    original_bytes = b"fake-mp4-payload-" + tag.encode()
    film_path.write_bytes(original_bytes)

    film = FilmOutput(
        project_id=pid, episode_id=None, version_no=1,
        url=f"/storage/outputs/{film_name}",
        resolution="720P", aspect_ratio="9:16", duration=10, is_current=True,
    )
    db.add(film)
    await db.commit()

    try:
        # ---- 不带 resolution：原片直下，字节一致 ----
        resp = await client.get(
            f"/api/publish/outputs/{film.id}/download", headers=auth["headers"]
        )
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith("video/mp4")
        assert resp.content == original_bytes

        # ---- 带 resolution=720p：有 ffmpeg 转码失败降级原片；无 ffmpeg 直下 ----
        resp2 = await client.get(
            f"/api/publish/outputs/{film.id}/download",
            params={"resolution": "720p"}, headers=auth["headers"],
        )
        assert resp2.status_code == 200, resp2.text
        assert resp2.headers["content-type"].startswith("video/mp4")
        # 伪 MP4 转码必然失败 → 降级返回原片字节（真有 ffmpeg 才可能转码成功）
        assert resp2.content == original_bytes

        # ---- 非法 resolution → 400 ----
        bad = await client.get(
            f"/api/publish/outputs/{film.id}/download",
            params={"resolution": "4k"}, headers=auth["headers"],
        )
        assert bad.status_code == 400

        # ---- 非本人成片 → 404 ----
        other = await register_user(client, prefix="other")
        other_headers = {"Authorization": f"Bearer {other['access_token']}"}
        denied = await client.get(
            f"/api/publish/outputs/{film.id}/download", headers=other_headers
        )
        assert denied.status_code == 404
    finally:
        film_path.unlink(missing_ok=True)
